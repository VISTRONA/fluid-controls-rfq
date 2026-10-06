"""Company tracker import/export. Original cells never stand in for workflow decisions."""
import io
import json
import zipfile
from xml.etree.ElementTree import ParseError
from openpyxl.utils.exceptions import InvalidFileException
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path
from flask import current_app
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from models import db, RFQ, User, StatusHistory, AuditLog

TRACKER_HEADERS = [
    'RFQ Received Date', 'FCL Enquiry No', 'Sales Person', 'Request From',
    'R&D Primary Person', 'R&D Review', 'Status', 'Completion Date',
    'Customer Name /Project Specification', 'Remark/Justification for Sales',
    'Requirement', 'GA Drawing Status', 'Request For Broughtout Parts Status',
    'Quotation Status For Brought Out Parts']
EXTRA_HEADERS = ['Workflow Status', 'Title', 'Priority', 'SLA Deadline',
                 'Quotation Deadline', 'Currency', 'Estimated Value',
                 'Current RFQ Type', 'Current Customer/Project', 'Current Assignee']


def normalize(value):
    return ' '.join(str(value or '').split()).casefold()


def cell_value(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if not value:
        return None
    # Only unambiguous ISO dates or explicitly day-first tracker dates.
    for fmt in ('%Y-%m-%d', '%Y-%m-%dT%H:%M:%S', '%d/%m/%Y', '%d-%m-%Y'):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    raise ValueError(f'Invalid date: {value}; use YYYY-MM-DD')


class SpreadsheetService:
    @staticmethod
    def preview(content):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                if sum(i.file_size for i in z.infolist()) > 30 * 1024 * 1024:
                    raise ValueError('Expanded workbook exceeds 30 MB')
            wb = load_workbook(io.BytesIO(content), data_only=False, read_only=True)
        except (zipfile.BadZipFile, OSError, KeyError, ParseError, InvalidFileException, EOFError, RuntimeError) as exc:
            raise ValueError('Upload a valid .xlsx workbook') from exc
        rows = []
        try:
            for sheet in wb:
                if sheet.max_row > 5001 or sheet.max_column > 100:
                    raise ValueError('Workbook limit: 5,000 rows and 100 columns per sheet')
                first = next(sheet.iter_rows(max_row=1), ())
                headers = [normalize(c.value) for c in first]
                if not all(normalize(h) in headers for h in TRACKER_HEADERS):
                    continue  # Display and Workng are not RFQ tables.
                if len(set(h for h in headers if h)) != len([h for h in headers if h]):
                    raise ValueError(f'Duplicate column names in {sheet.title}')
                indices = {h: headers.index(normalize(h)) for h in TRACKER_HEADERS + EXTRA_HEADERS if normalize(h) in headers}
                for cells in sheet.iter_rows(min_row=2):
                    if not any(c.value is not None for c in cells):
                        continue
                    raw = {h: cell_value(cells[i].value) for h, i in indices.items()}
                    errors = []
                    if any(c.data_type == 'f' for c in cells):
                        errors.append('Formula cells are not imported; replace with reviewed values')
                    number = str(raw.get('FCL Enquiry No') or '').strip()
                    if not number:
                        errors.append('Enquiry number is required')
                    if len(number) > 50:
                        errors.append('Enquiry number exceeds 50 characters')
                    rows.append({'sheet': sheet.title, 'row': cells[0].row, 'raw': raw,
                                 'number': number, 'errors': errors})
        except (ParseError, zipfile.BadZipFile, EOFError, KeyError) as exc:
            raise ValueError('Workbook contents could not be read; upload a valid .xlsx') from exc
        finally:
            wb.close()
        if len(rows) > 5000:
            raise ValueError('Workbook limit: 5,000 total RFQ rows')
        if not rows:
            raise ValueError('No RFQ sheet found with all 14 company tracker headers in row 1')
        numbers = Counter(r['number'] for r in rows)
        existing = {n for (n,) in db.session.query(RFQ.rfq_number).all()}
        for row in rows:
            if numbers[row['number']] > 1:
                row['warnings'] = ['Duplicate enquiry in workbook: select at most one after review']
            else:
                row['warnings'] = []
            raw = row['raw']
            if not raw.get('Customer Name /Project Specification') or raw.get('Customer Name /Project Specification') == '-':
                row['warnings'].append('Customer/project is missing; correct it before selecting this row')
            if not raw.get('Status'):
                row['warnings'].append('Historical status is blank; an explicit workflow choice is required')
            for header in ('RFQ Received Date', 'Completion Date', 'SLA Deadline', 'Quotation Deadline'):
                try:
                    if parse_date(raw.get(header)) is None and header == 'RFQ Received Date':
                        row['warnings'].append('Received date is missing; enter a reviewed date')
                except ValueError:
                    row['warnings'].append(f'{header} is invalid; enter a reviewed date')
            if row['number'] in existing:
                row['errors'].append('Enquiry already exists; import creates new records only')
        return rows

    @staticmethod
    def prepare(rows, form):
        """Revalidate selected rows at commit; any error prevents the whole batch."""
        c = current_app.config
        selected = set(form.getlist('include'))
        records, errors = [], []
        seen = set()
        statuses = sorted({str(r['raw'].get('Status') or '') for r in rows})
        types = sorted({str(r['raw'].get('Request From') or '') for r in rows})
        sales = sorted({str(r['raw'].get('Sales Person') or '') for r in rows})
        maps = {s: form.get(f'status_{i}', '') for i, s in enumerate(statuses)}
        type_maps = {s: form.get(f'type_{i}', '') for i, s in enumerate(types)}
        sales_maps = {s: form.get(f'sales_{i}', '') for i, s in enumerate(sales)}
        for idx, row in enumerate(rows):
            if str(idx) not in selected:
                continue
            raw = row['raw']
            issues = list(row['errors'])
            number = row['number']
            if number in seen or RFQ.query.filter_by(rfq_number=number).first():
                issues.append('Duplicate/existing enquiry number')
            seen.add(number)
            status = maps[str(raw.get('Status') or '')]
            kind = type_maps[str(raw.get('Request From') or '')]
            if status not in c['STATUSES']:
                issues.append('Choose an explicit historical status mapping')
            if kind not in c['RFQ_TYPES']:
                issues.append('Choose an explicit Request From mapping')
            assigned = sales_maps[str(raw.get('Sales Person') or '')]
            user = db.session.get(User, int(assigned)) if assigned.isdigit() else None
            if assigned != 'unassigned' and (not user or not user.is_active):
                issues.append('Choose a Sales Person account or explicitly leave unassigned')
            customer = form.get(f'customer_{idx}', '').strip()
            title = form.get(f'title_{idx}', '').strip()
            if not customer or customer == '-' or len(customer) > 150:
                issues.append('Customer/project required, maximum 150 characters')
            if not title or len(title) > 200:
                issues.append('Title required, maximum 200 characters')
            priority = form.get(f'priority_{idx}', '')
            if priority not in c['PRIORITIES']:
                issues.append('Choose priority')
            try:
                received = parse_date(form.get(f'received_{idx}'))
                completed = parse_date(form.get(f'completed_{idx}'))
                sla = parse_date(form.get(f'sla_{idx}'))
                quote = parse_date(form.get(f'quote_{idx}'))
                if received is None:
                    raise ValueError('Received date required')
                if sla is None or quote is None:
                    if form.get('apply_sla_defaults') != 'yes':
                        raise ValueError('Enter both deadlines or explicitly approve existing priority/calendar-day defaults')
                    sla = sla or received + timedelta(days=c['SLA_DAYS_BY_PRIORITY'].get(priority, 7))
                    quote = quote or sla
                if sla < received or quote < received:
                    raise ValueError('Deadlines cannot precede received date')
                if status in c['TERMINAL_STATUSES'] and completed is None:
                    raise ValueError('Terminal workflow status requires reviewed completion date')
                if status not in c['TERMINAL_STATUSES'] and completed:
                    raise ValueError('Completion date conflicts with active workflow status; review or clear it')
                if completed and completed < received:
                    raise ValueError('Completion precedes received date')
            except ValueError as exc:
                issues.append(str(exc))
            currency = str(raw.get('Currency') or 'INR')
            try:
                value = float(raw.get('Estimated Value') or 0)
                import math
                if not math.isfinite(value) or value < 0 or currency not in c['CURRENCIES']:
                    raise ValueError()
            except (ValueError, TypeError):
                issues.append('Invalid estimated value/currency')
            if issues:
                errors.append(f'{row["sheet"]} row {row["row"]} ({number}): ' + '; '.join(issues))
                continue
            records.append(dict(rfq_number=number, title=title, customer_name=customer,
                rfq_type=kind, received_date=received, sla_deadline=sla,
                quotation_deadline=quote, current_status=status, priority=priority,
                assigned_to_id=user.id if user else None, estimated_value=value, currency=currency,
                completed_at=datetime.combine(completed, datetime.min.time()) if completed else None,
                remarks=str(raw.get('Remark/Justification for Sales') or ''),
                tracker_data={'sheet': row['sheet'], 'row': row['row'], 'cells': raw}))
        if not selected:
            errors.append('Select at least one reviewed row')
        if selected - {str(i) for i in range(len(rows))}:
            errors.append('Invalid row selection')
        if form.get('review_confirmed') != 'yes':
            errors.append('Confirm the selected rows and mappings were reviewed')
        return records, errors

    @staticmethod
    def commit(records, user, ip):
        for data in records:
            rfq = RFQ(**data, created_by_id=user.id)
            db.session.add(rfq)
            db.session.flush()
            db.session.add(StatusHistory(rfq_id=rfq.id, to_status=rfq.current_status,
                changed_by_id=user.id, comments=f'Reviewed tracker import; original status: {data["tracker_data"]["cells"].get("Status")!r}'))
            db.session.add(AuditLog(user_id=user.id, rfq_id=rfq.id, action='TRACKER_IMPORTED',
                description=f'Imported {rfq.rfq_number}; source {data["tracker_data"]["sheet"]} row {data["tracker_data"]["row"]}', ip_address=ip))
        # One transaction; historical imports intentionally create no completion/assignment emails.
        db.session.commit()

    @staticmethod
    def export(rfqs):
        wb = Workbook()
        tracker = wb.active
        tracker.title = 'RFQ Tracker'
        tracker.append(TRACKER_HEADERS + EXTRA_HEADERS)
        data_sheet = wb.create_sheet('RFQ Data')
        columns = [col.name for col in RFQ.__table__.columns if col.name != 'tracker_data']
        data_sheet.append(columns)
        quotes = wb.create_sheet('Quotations')
        quotes.append(['FCL Enquiry No', 'Quotation Number', 'Date', 'Amount', 'Currency', 'Validity', 'Status', 'Submitted Date', 'Remarks'])
        history = wb.create_sheet('Status History')
        history.append(['FCL Enquiry No', 'From Status', 'To Status', 'Date', 'Comments'])
        followups = wb.create_sheet('Follow-ups')
        followups.append(['FCL Enquiry No', 'Follow-up Date', 'Contact Person', 'Method', 'Status', 'Next Follow-up', 'Comments'])
        for r in rfqs:
            raw = (r.tracker_data or {}).get('cells', {})
            values = [raw.get(h) for h in TRACKER_HEADERS]
            # Imported historical cells stay intact. Current workflow is an explicit extra column.
            if not raw:
                values = [r.received_date, r.rfq_number, r.assigned_employee.name if r.assigned_employee else '',
                    r.rfq_type, '', '', r.current_status, r.completed_at.date() if r.completed_at else None,
                    r.customer_name, r.remarks, r.product_service, '', '', '']
            for idx in (0, 7):
                try:
                    values[idx] = parse_date(values[idx])
                except ValueError:
                    pass
            tracker.append(values + [r.current_status, r.title, r.priority, r.sla_deadline,
                                      r.quotation_deadline, r.currency, r.estimated_value, r.rfq_type, r.customer_name,
                                      r.assigned_employee.name if r.assigned_employee else 'Unassigned'])
            data_sheet.append([getattr(r, col) for col in columns])
            for q in r.quotations:
                quotes.append([r.rfq_number, q.quotation_number, q.quotation_date, q.quotation_amount,
                    q.currency, q.validity, q.status, q.submitted_date, q.remarks])
            for h in r.status_history:
                history.append([r.rfq_number, h.from_status, h.to_status, h.created_at, h.comments])
            for f in r.followups:
                followups.append([r.rfq_number, f.followup_date, f.contact_person, f.contact_method,
                    f.status, f.next_followup_date, f.comments])
        notes = wb.create_sheet('Read Me')
        notes.append(['Fluid Controls RFQ export'])
        notes.append(['Tracker Status and original tracker cells preserve historical source values; Workflow Status is the current application status.'])
        notes.append(['RFQ Data contains current database fields. No currency conversion or status equivalence is assumed.'])
        notes.append(['Tracker import creates new RFQs only, requires reviewed mappings and deadlines, and is not a full database restore.'])
        for sheet in wb:
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.orientation = 'landscape'
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.print_title_rows = '1:1'
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='234968')
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    # Write all user text as literal strings, never executable Excel formulas.
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
                    if isinstance(cell.value, (date, datetime)):
                        cell.number_format = 'yyyy-mm-dd'
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
            for col in sheet.columns:
                sheet.column_dimensions[col[0].column_letter].width = 24
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
