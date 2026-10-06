"""Focused tests use disposable SQLite databases; no company workbook or seed data needed."""
import io
import os
import re
import sqlite3
from datetime import date, datetime, timedelta
from unittest.mock import patch, MagicMock

os.environ['DATABASE_URL'] = 'sqlite:///:memory:'
import pytest
from werkzeug.datastructures import MultiDict
from openpyxl import Workbook, load_workbook
from app import create_app
from config import Config
from models import db, RFQ, User, Notification, Quotation, StatusHistory
from services.rfq_service import RFQService
from services.notification_service import NotificationService
from services.spreadsheet_service import SpreadsheetService, TRACKER_HEADERS


@pytest.fixture
def app(tmp_path):
    class TestConfig(Config):
        TESTING = True
        SECRET_KEY = 'test-secret'
        SQLALCHEMY_DATABASE_URI = 'sqlite:///:memory:'
        EMAIL_MODE = 'disabled'
        MAJOR_RFQ_PRIORITIES = ['High']
    application = create_app(TestConfig)
    application.instance_path = str(tmp_path)
    @application.before_request
    def reset_test_login_cache():
        from flask import g
        g.pop('_login_user', None)
    with application.app_context():
        for name, role in [('admin', 'ADMIN'), ('manager', 'MANAGER'), ('sales', 'USER'), ('other', 'USER')]:
            user = User(username=name, email=f'{name}@example.test', name=name, role=role)
            user.set_password('test-password-long')
            db.session.add(user)
        db.session.commit()
        yield application
        db.session.remove()
        db.drop_all()


@pytest.fixture
def client(app):
    c = app.test_client()
    with c.session_transaction() as s:
        s['_user_id'] = '1'
        s['_fresh'] = True
        s['_csrf'] = 'test-csrf'
    return c


def make_rfq(number='FC001', days=1, priority='High', status='RECEIVED'):
    r = RFQ(rfq_number=number, title='Fitting', customer_name='Example Customer',
        received_date=date.today() - timedelta(days=5), sla_deadline=date.today() + timedelta(days=days),
        quotation_deadline=date.today() + timedelta(days=days), priority=priority,
        current_status=status, assigned_to_id=3, currency='INR')
    db.session.add(r)
    db.session.commit()
    return r


def workbook(rows=None):
    w = Workbook()
    w.active.title = 'RFQ 2025-2026'
    w.active.append([h + ' ' if h == 'Request From' else h for h in TRACKER_HEADERS])
    for row in rows or [[datetime(2026, 4, 2), 'FC123', 'Malini', 'Railways', 'Divya', 'Krishna',
                            'Complete', datetime(2026, 4, 3), 'Customer', 'Original remark', 'FIRM', 'Yes', 'Yes', 'Not Received']]:
        w.active.append(row)
    w.create_sheet('Display').append(['RFQ'])
    out = io.BytesIO()
    w.save(out)
    return out.getvalue()


def form_for(rows):
    form = MultiDict({'include': '0', 'status_0': 'CLOSED', 'type_0': 'Railway',
        'sales_0': 'unassigned', 'customer_0': 'Customer', 'title_0': 'FC123',
        'priority_0': 'Medium', 'received_0': '2026-04-02', 'completed_0': '2026-04-03',
        'sla_0': '', 'quote_0': '', 'apply_sla_defaults': 'yes', 'review_confirmed': 'yes'})
    return form


def test_sla_dedup_read_and_new_deadline(app):
    r = make_rfq()
    NotificationService.scan()
    assert Notification.query.count() == 3  # assignee plus active managers/admin
    Notification.query.update({'is_read': True})
    db.session.commit()
    NotificationService.scan()
    assert Notification.query.count() == 3
    r.sla_deadline += timedelta(days=1)
    db.session.commit()
    NotificationService.scan()
    assert Notification.query.count() == 6
    assert {n.user_id for n in Notification.query.all()} == {1, 2, 3}


def test_breach_completion_and_reopen(app):
    r = make_rfq(days=-1)
    NotificationService.scan()
    assert {n.notification_type for n in Notification.query.all()} == {'sla_breached'}
    RFQService.update_status(r.id, 'WON', 'Done', db.session.get(User, 1))
    assert Notification.query.filter_by(notification_type='rfq_completed').count() == 3
    NotificationService.scan()
    assert Notification.query.filter_by(notification_type='sla_breached').count() == 3
    RFQService.update_status(r.id, 'WON', '', db.session.get(User, 1))
    RFQService.update_status(r.id, 'CLOSED', '', db.session.get(User, 1))
    assert Notification.query.filter_by(notification_type='rfq_completed').count() == 3
    RFQService.update_status(r.id, 'RECEIVED', '', db.session.get(User, 1))
    RFQService.update_status(r.id, 'CLOSED', '', db.session.get(User, 1))
    assert Notification.query.filter_by(notification_type='rfq_completed').count() == 6
    with pytest.raises(ValueError):
        RFQService.update_status(r.id, 'Complete', '', db.session.get(User, 1))


def test_no_sla_for_terminal_or_far_deadline(app):
    make_rfq(days=30)
    make_rfq(number='FC002', status='WON')
    NotificationService.scan()
    assert Notification.query.count() == 0


def test_email_disabled_log_and_major_selection(app):
    r = make_rfq()
    NotificationService.scan()
    with patch('services.notification_service.smtplib.SMTP') as smtp:
        NotificationService.deliver()
        smtp.assert_not_called()
    assert all(n.email_status == 'disabled' for n in Notification.query.all())
    app.config['EMAIL_MODE'] = 'log'
    minor = make_rfq(number='FC002', priority='Low', days=30)
    RFQService.update_status(minor.id, 'CLOSED', '', db.session.get(User, 1))
    assert all(n.email_status == 'disabled' for n in Notification.query.filter_by(notification_type='rfq_completed'))
    RFQService.update_status(r.id, 'WON', '', db.session.get(User, 1))
    NotificationService.deliver()
    assert all(n.email_status == 'logged' for n in Notification.query.filter_by(notification_type='rfq_completed', rfq_id=r.id))


def test_smtp_tls_and_failed_retry(app):
    app.config.update(EMAIL_MODE='smtp', SMTP_HOST='smtp.test', SMTP_FROM='rfq@example.test',
        SMTP_USERNAME='user', SMTP_PASSWORD='secret')
    make_rfq()
    NotificationService.scan()
    with patch('services.notification_service.smtplib.SMTP') as smtp:
        connection = smtp.return_value.__enter__.return_value
        NotificationService.deliver()
        assert connection.starttls.call_count == 3
        assert connection.login.call_count == 3
        assert connection.send_message.call_count == 3
        NotificationService.deliver()
        assert connection.send_message.call_count == 3
    assert all(n.email_status == 'sent' for n in Notification.query.all())
    make_rfq(number='FC002')
    NotificationService.scan()
    with patch('services.notification_service.smtplib.SMTP', side_effect=OSError('failed')) as smtp:
        NotificationService.deliver()
        NotificationService.deliver()
        assert smtp.call_count == 3
    assert all(n.email_attempts == 1 for n in Notification.query.filter_by(email_status='pending'))


def test_stale_sla_email_skipped(app):
    app.config['EMAIL_MODE'] = 'log'
    r = make_rfq()
    NotificationService.scan()
    r.sla_deadline += timedelta(days=20)
    db.session.commit()
    NotificationService.deliver()
    assert all(n.email_status == 'skipped' for n in Notification.query.all())


def test_import_requires_explicit_mapping_and_sla_approval(app):
    rows = SpreadsheetService.preview(workbook())
    assert len(rows) == 1
    assert rows[0]['raw']['Status'] == 'Complete'
    form = form_for(rows)
    form['status_0'] = ''
    _, errors = SpreadsheetService.prepare(rows, form)
    assert any('explicit historical' in e for e in errors)
    form['status_0'] = 'CLOSED'
    del form['apply_sla_defaults']
    _, errors = SpreadsheetService.prepare(rows, form)
    assert any('approve' in e for e in errors)
    assert RFQ.query.count() == 0


def test_import_commit_preserves_source_no_historical_alerts(app):
    app.config['EMAIL_MODE'] = 'log'
    rows = SpreadsheetService.preview(workbook())
    records, errors = SpreadsheetService.prepare(rows, form_for(rows))
    assert errors == []
    SpreadsheetService.commit(records, db.session.get(User, 1), '127.0.0.1')
    r = RFQ.query.one()
    assert r.current_status == 'CLOSED' and r.rfq_type == 'Railway'
    assert r.tracker_data['cells']['Status'] == 'Complete'
    assert r.tracker_data['cells']['R&D Primary Person'] == 'Divya'
    assert r.completed_at.date() == date(2026, 4, 3)
    assert r.sla_deadline == date(2026, 4, 9)
    NotificationService.scan()
    assert Notification.query.count() == 0
    assert StatusHistory.query.one().to_status == 'CLOSED'


def test_duplicate_rows_and_existing_record_block_import(app):
    w = load_workbook(io.BytesIO(workbook()))
    w.active.append([c.value for c in w.active[2]])
    out = io.BytesIO(); w.save(out)
    rows = SpreadsheetService.preview(out.getvalue())
    assert rows[0]['warnings']
    f = form_for(rows)
    f.add('include', '1')
    for key in ['customer', 'title', 'priority', 'received', 'completed', 'sla', 'quote']:
        f[f'{key}_1'] = f[f'{key}_0']
    _, errors = SpreadsheetService.prepare(rows, f)
    assert any('Duplicate' in e for e in errors)
    f.setlist('include', ['0'])
    records, errors = SpreadsheetService.prepare(rows, f)
    assert not errors
    SpreadsheetService.commit(records, db.session.get(User, 1), None)
    assert 'already exists' in SpreadsheetService.preview(workbook())[0]['errors'][0]


@pytest.mark.parametrize('field,value,expected', [
    ('received_0', '04/02/26', 'Invalid date'),
    ('completed_0', '', 'completion date'),
    ('completed_0', '2026-04-01', 'precedes'),
    ('sla_0', '2026-04-01', 'Deadlines'),
    ('customer_0', '-', 'Customer'),
    ('status_0', 'RECEIVED', 'conflicts'),
    ('sales_0', '999', 'Sales Person'),
    ('review_confirmed', '', 'Confirm')])
def test_import_validation(app, field, value, expected):
    rows = SpreadsheetService.preview(workbook())
    f = form_for(rows); f[field] = value
    _, errors = SpreadsheetService.prepare(rows, f)
    assert any(expected in e for e in errors)


def test_formula_workbook_rejected_and_export_text_safe(app):
    w = load_workbook(io.BytesIO(workbook()))
    w.active['I2'] = '=HYPERLINK("http://example.test","Click")'
    out = io.BytesIO(); w.save(out)
    rows = SpreadsheetService.preview(out.getvalue())
    assert any('Formula' in e for e in rows[0]['errors'])
    r = make_rfq(); r.title = '=1+1'; db.session.commit()
    exported = load_workbook(SpreadsheetService.export([r]), data_only=False)
    assert exported['RFQ Tracker']['P2'].value == '=1+1'
    assert exported['RFQ Tracker']['P2'].data_type == 's'
    assert exported['RFQ Tracker']['A2'].is_date
    assert exported['RFQ Tracker'].freeze_panes == 'A2'
    assert set(exported.sheetnames) == {'RFQ Tracker', 'RFQ Data', 'Quotations', 'Status History', 'Follow-ups', 'Read Me'}


def test_export_original_and_current_status_distinct(app):
    r = make_rfq()
    r.tracker_data = {'sheet': 'Original', 'row': 2, 'cells': {'FCL Enquiry No': r.rfq_number,
        'Status': 'In-Process', 'R&D Review': 'Krishna', 'RFQ Received Date': '2026-04-02'}}
    db.session.commit()
    w = load_workbook(SpreadsheetService.export([r]))
    assert w['RFQ Tracker']['G2'].value == 'In-Process'
    assert w['RFQ Tracker']['O2'].value == 'RECEIVED'
    assert w['RFQ Tracker']['F2'].value == 'Krishna'


def test_page_endpoints_and_auth(client, app):
    r = make_rfq()
    q = Quotation(rfq_id=r.id, quotation_number='Q1', quotation_amount=125, currency='INR')
    db.session.add(q); db.session.commit()
    for endpoint in ['/dashboard', '/rfqs', f'/rfqs/{r.id}', '/reports', '/rfqs/import',
        '/rfqs/print', f'/rfqs/{r.id}/print', '/reports/print',
        f'/rfqs/{r.id}/quotations/{q.id}/print', '/notifications', '/api/rfqs']:
        assert client.get(endpoint).status_code == 200, endpoint
    assert client.get(f'/rfqs/{r.id}/quotations/999/print').status_code == 404
    anon = app.test_client()
    assert anon.get('/reports/export.xlsx').status_code == 302
    assert anon.get('/rfqs/import').status_code == 302


def test_filter_export_print_all_pages(client, app):
    for i in range(15):
        make_rfq(number=f'FILTER{i}', days=30)
    other = make_rfq(number='OTHER', days=30); other.customer_name='Other'; db.session.commit()
    response = client.get('/reports/export.xlsx?search=FILTER&status=RECEIVED&page=2')
    w = load_workbook(io.BytesIO(response.data))
    assert w['RFQ Tracker'].max_row == 16
    printed = client.get('/rfqs/print?search=FILTER')
    assert b'15 RFQs' in printed.data and b'OTHER' not in printed.data
    assert b'FILTER14' in client.get('/rfqs?search=FILTER').data


def test_notification_ownership_and_csrf(client, app):
    make_rfq()
    response = client.get('/notifications/unread-count')
    assert response.json['count'] == 1
    own = Notification.query.filter_by(user_id=1).one()
    other = Notification.query.filter_by(user_id=3).one()
    assert client.get(f'/notifications/mark-read/{own.id}').status_code == 405
    assert client.post(f'/notifications/mark-read/{own.id}').status_code == 400
    assert client.post(f'/notifications/mark-read/{other.id}', data={'csrf_token':'test-csrf'}).status_code == 404
    assert client.post(f'/notifications/mark-read/{own.id}', data={'csrf_token':'test-csrf'}).status_code == 302
    assert client.get('/notifications/unread-count').json['count'] == 0


def test_import_endpoint_replay_role_and_atomic_validation(client, app):
    response = client.post('/rfqs/import', data={'csrf_token':'test-csrf','action':'preview',
        'workbook':(io.BytesIO(workbook()), 'tracker.xlsx')})
    assert response.status_code == 200 and b'Complete' in response.data
    assert RFQ.query.count() == 0
    with client.session_transaction() as s:
        token = s['tracker_import_token']
    f = form_for([])
    f.update({'csrf_token':'test-csrf','action':'commit','preview_token':token})
    f['status_0'] = ''
    assert client.post('/rfqs/import', data=f).status_code == 200
    assert RFQ.query.count() == 0
    f['status_0'] = 'CLOSED'
    assert client.post('/rfqs/import', data=f).status_code == 302
    assert RFQ.query.count() == 1
    assert b'Preview expired' in client.post('/rfqs/import', data=f).data
    with client.session_transaction() as s:
        s['_user_id']='4'
    assert client.get('/rfqs/import').status_code == 302
    assert client.post('/rfqs/1/status',data={'csrf_token':'test-csrf','new_status':'WON'}).status_code == 403


def test_delete_cascades_notifications(client, app):
    r = make_rfq()
    NotificationService.scan()
    response = client.delete(f'/api/rfqs/{r.id}')
    assert response.status_code == 200
    assert Notification.query.count() == 0


def test_legacy_database_upgrade_is_idempotent(tmp_path):
    path = tmp_path/'old.db'
    with sqlite3.connect(path) as conn:
        conn.execute('CREATE TABLE rfqs (id INTEGER PRIMARY KEY, rfq_number TEXT)')
        conn.execute("INSERT INTO rfqs VALUES (1,'PRESERVED')")
        conn.execute('CREATE TABLE notifications (id INTEGER PRIMARY KEY)')
    class LegacyConfig(Config):
        SQLALCHEMY_DATABASE_URI = f'sqlite:///{path}'
        TESTING = True
    a = create_app(LegacyConfig)
    runner = a.test_cli_runner()
    for _ in range(2):
        result = runner.invoke(args=['upgrade-db'])
        assert result.exit_code == 0, result.output
    with sqlite3.connect(path) as conn:
        assert conn.execute('SELECT rfq_number FROM rfqs').fetchone()[0] == 'PRESERVED'
        assert 'tracker_data' in {r[1] for r in conn.execute('PRAGMA table_info(rfqs)')}
        assert 'email_status' in {r[1] for r in conn.execute('PRAGMA table_info(notifications)')}


def test_email_retry_limit_and_inactive_recipient(app):
    app.config['EMAIL_MODE'] = 'smtp'
    app.config.update(SMTP_HOST='smtp.test', SMTP_FROM='rfq@example.test')
    make_rfq()
    NotificationService.scan()
    db.session.get(User, 3).is_active = False
    db.session.commit()
    with patch('services.notification_service.smtplib.SMTP', side_effect=OSError('unavailable')) as smtp:
        for _ in range(6):
            NotificationService.deliver()
            for n in Notification.query.filter_by(email_status='pending'):
                n.email_attempted_at = datetime.utcnow() - timedelta(seconds=1000)
            db.session.commit()
        assert smtp.call_count == 10
    assert Notification.query.filter_by(email_status='failed').count() == 2
    assert Notification.query.filter_by(email_status='skipped').count() == 1


def test_api_creation_and_failed_duplicate_rollback(client, app):
    payload = {'rfq_number':'API1','title':'Fitting','customer_name':'Customer',
        'received_date':'2026-04-02','sla_deadline':'2026-04-09','assigned_to_id':3}
    response = client.post('/api/rfqs', json=payload)
    assert response.status_code == 201
    assert response.json['rfq']['current_status'] == 'RECEIVED'
    assert Notification.query.filter_by(notification_type='assigned').count() == 1
    assert client.post('/api/rfqs',json=payload).status_code == 400
    assert client.get('/api/rfqs').status_code == 200
    assert RFQ.query.count() == 1


def test_multirow_commit_rolls_back_on_conflict(app):
    rows = SpreadsheetService.preview(workbook())
    records, errors = SpreadsheetService.prepare(rows, form_for(rows))
    assert not errors
    conflict = dict(records[0])
    from sqlalchemy.exc import IntegrityError
    with pytest.raises(IntegrityError):
        SpreadsheetService.commit(records + [conflict], db.session.get(User, 1), None)
    db.session.rollback()
    assert RFQ.query.count() == 0
    assert StatusHistory.query.count() == 0
